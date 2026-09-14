import csv
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


def render(
    csv_name: str,
    xlsx_name: str,
    sheet_name: str,
    title: str,
):
    """
    Читает CSV и создаёт оформленный XLSX.

    Этот файл является отдельным приложением,
    которое запускает основное приложение.
    """
    with open(
        csv_name,
        encoding="utf-8-sig",
        newline="",
    ) as file:
        rows = list(csv.reader(file))

    book = Workbook()
    sheet = book.active
    sheet.title = sheet_name[:31]

    for row in rows:
        sheet.append(row)

    last_column = max(1, sheet.max_column)

    sheet.insert_rows(1)

    sheet.merge_cells(
        start_row=1,
        start_column=1,
        end_row=1,
        end_column=last_column,
    )

    title_cell = sheet.cell(1, 1, title)
    title_cell.font = Font(
        bold=True,
        color="FFFFFF",
        size=14,
    )
    title_cell.fill = PatternFill(
        "solid",
        fgColor="17365D",
    )
    title_cell.alignment = Alignment(
        horizontal="center",
        vertical="center",
    )

    header_row = 2

    header_fill = PatternFill(
        "solid",
        fgColor="1F4E78",
    )

    thin_border = Side(
        style="thin",
        color="D9E2F3",
    )

    for cell in sheet[header_row]:
        cell.font = Font(
            bold=True,
            color="FFFFFF",
        )
        cell.fill = header_fill
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )
        cell.border = Border(bottom=thin_border)

    for row in range(header_row + 1, sheet.max_row + 1):
        for column in range(1, sheet.max_column + 1):
            cell = sheet.cell(row, column)

            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )

            cell.border = Border(bottom=thin_border)

            if row % 2 == 0:
                cell.fill = PatternFill(
                    "solid",
                    fgColor="EAF2F8",
                )

            header = str(
                sheet.cell(header_row, column).value or ""
            )

            if "₽" in header:
                cell.number_format = '#,##0.00 [$₽-419]'

    for column in range(1, sheet.max_column + 1):
        values = [
            str(sheet.cell(row, column).value or "")
            for row in range(1, sheet.max_row + 1)
        ]

        width = min(
            max(max(map(len, values)) + 2, 12),
            42,
        )

        sheet.column_dimensions[
            get_column_letter(column)
        ].width = width

    sheet.freeze_panes = "A3"

    sheet.auto_filter.ref = (
        f"A{header_row}:"
        f"{get_column_letter(sheet.max_column)}"
        f"{sheet.max_row}"
    )

    target = Path(xlsx_name)
    target.parent.mkdir(exist_ok=True)

    book.save(target)


if __name__ == "__main__":
    if len(sys.argv) != 5:
        raise SystemExit(
            "Использование: render_report.py "
            "input.csv output.xlsx sheet_name title"
        )

    render(
        sys.argv[1],
        sys.argv[2],
        sys.argv[3],
        sys.argv[4],
    )
