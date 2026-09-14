from pathlib import Path
import csv
import sqlite3
import subprocess
import sys

import psycopg
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .config import pg_dsn
from .importer import SOURCE

REPORTS = Path("reports")
RENDER_SCRIPT = Path("scripts/render_report.py")


def write_csv(path: Path, headers, rows):
    """Записывает набор строк в CSV в UTF-8 с BOM для корректного Excel."""
    path.parent.mkdir(exist_ok=True)

    with path.open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.writer(file)
        writer.writerow(headers)
        writer.writerows(rows)


def render_xlsx(
    csv_path: Path,
    xlsx_path: Path,
    sheet_name: str,
    title: str,
):
    """Запускает отдельное приложение, которое создаёт XLSX из CSV."""
    subprocess.run(
        [
            sys.executable,
            str(RENDER_SCRIPT),
            str(csv_path),
            str(xlsx_path),
            sheet_name,
            title,
        ],
        check=True,
    )


def export_source_files():
    """
    Экспортирует исходную ненормализованную SQLite-таблицу.

    Результат:
    reports/01_source_denormalized.csv
    reports/01_source_denormalized.xlsx
    """
    if not SOURCE.exists():
        raise FileNotFoundError(
            "Нет SQLite-файла. Сначала выполните create-source."
        )

    with sqlite3.connect(SOURCE) as conn:
        cursor = conn.execute(
            "SELECT * FROM sales_flat ORDER BY sale_number"
        )
        headers = [column[0] for column in cursor.description]
        rows = cursor.fetchall()

    csv_path = REPORTS / "01_source_denormalized.csv"
    xlsx_path = REPORTS / "01_source_denormalized.xlsx"

    write_csv(csv_path, headers, rows)

    render_xlsx(
        csv_path,
        xlsx_path,
        "Исходные данные",
        "Шаг 1. Ненормализованные данные SQLite",
    )

    return len(rows), csv_path, xlsx_path


NORMALIZED_QUERY = """
SELECT
    s.sale_number AS "Номер продажи",
    s.sold_at AS "Дата продажи",
    c.full_name AS "Покупатель",
    c.phone AS "Телефон",
    e.full_name AS "Сотрудник",
    b.isbn AS "ISBN",
    b.title AS "Книга",
    g.name AS "Жанр",
    STRING_AGG(a.full_name, ', ' ORDER BY a.full_name) AS "Авторы",
    s.quantity AS "Количество",
    s.unit_price AS "Цена, ₽",
    s.quantity * s.unit_price AS "Сумма, ₽"
FROM sales s
JOIN customers c ON c.customer_id = s.customer_id
JOIN employees e ON e.employee_id = s.employee_id
JOIN books b ON b.book_id = s.book_id
JOIN genres g ON g.genre_id = b.genre_id
JOIN book_authors ba ON ba.book_id = b.book_id
JOIN authors a ON a.author_id = ba.author_id
GROUP BY
    s.sale_id,
    c.customer_id,
    e.employee_id,
    b.book_id,
    g.genre_id
ORDER BY s.sale_number
"""


def export_normalized_files():
    """
    Создаёт витрину нормализованных данных.

    В файле данные снова собраны в одну таблицу через JOIN.
    Это позволяет показать, что значения после нормализации не изменились.
    """
    with psycopg.connect(pg_dsn()) as conn:
        with conn.cursor() as cur:
            cur.execute(NORMALIZED_QUERY)
            headers = [column.name for column in cur.description]
            rows = cur.fetchall()

    csv_path = REPORTS / "02_normalized_data.csv"
    xlsx_path = REPORTS / "02_normalized_data.xlsx"

    write_csv(csv_path, headers, rows)

    render_xlsx(
        csv_path,
        xlsx_path,
        "Нормализованная витрина",
        "Шаг 2. Витрина из нормализованных таблиц PostgreSQL",
    )

    return len(rows), csv_path, xlsx_path


def apply_sheet_style(sheet, title: str):
    """
    Оформляет один лист Excel.

    Используется только для многостраничного файла
    02_normalized_structure.xlsx.
    """
    last_column = max(1, sheet.max_column)

    sheet.insert_rows(1)
    sheet.merge_cells(
        start_row=1,
        start_column=1,
        end_row=1,
        end_column=last_column,
    )

    title_cell = sheet.cell(1, 1, title)
    title_cell.font = Font(
        bold=True,
        color="FFFFFF",
        size=14,
    )
    title_cell.fill = PatternFill(
        "solid",
        fgColor="17365D",
    )
    title_cell.alignment = Alignment(
        horizontal="center",
        vertical="center",
    )

    header_row = 2
    header_fill = PatternFill(
        "solid",
        fgColor="1F4E78",
    )
    thin_border = Side(
        style="thin",
        color="D9E2F3",
    )

    for cell in sheet[header_row]:
        cell.font = Font(
            bold=True,
            color="FFFFFF",
        )
        cell.fill = header_fill
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )
        cell.border = Border(bottom=thin_border)

    for row in range(header_row + 1, sheet.max_row + 1):
        for column in range(1, sheet.max_column + 1):
            cell = sheet.cell(row, column)

            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )
            cell.border = Border(bottom=thin_border)

            if row % 2 == 0:
                cell.fill = PatternFill(
                    "solid",
                    fgColor="EAF2F8",
                )

            header = str(
                sheet.cell(header_row, column).value or ""
            )

            if "₽" in header:
                cell.number_format = '#,##0.00 [$₽-419]'

    for column in range(1, sheet.max_column + 1):
        values = [
            str(sheet.cell(row, column).value or "")
            for row in range(1, sheet.max_row + 1)
        ]

        width = min(
            max(max(map(len, values)) + 2, 12),
            42,
        )

        sheet.column_dimensions[
            get_column_letter(column)
        ].width = width

    sheet.freeze_panes = "A3"

    sheet.auto_filter.ref = (
        f"A{header_row}:"
        f"{get_column_letter(sheet.max_column)}"
        f"{sheet.max_row}"
    )


def add_sheet(book: Workbook, name: str, headers, rows, title: str):
    """Создаёт лист Excel, заполняет и оформляет его."""
    sheet = book.create_sheet(name[:31])

    sheet.append(headers)

    for row in rows:
        sheet.append(row)

    apply_sheet_style(sheet, title)


def export_normalized_structure():
    """
    Создаёт основной файл для объяснения нормализации.

    Внутри находятся:
    - Лист 'Сравнение'
    - Все 7 таблиц PostgreSQL по отдельности

    Это позволяет увидеть, что данные в PostgreSQL не хранятся
    одной повторяющейся таблицей.
    """
    comparison_headers = [
        "Критерий",
        "Ненормализованная SQLite",
        "Нормализованная PostgreSQL",
    ]

    comparison_rows = [
        [
            "Число таблиц",
            "1 таблица: sales_flat",
            (
                "7 таблиц: customers, employees, genres, "
                "authors, books, book_authors, sales"
            ),
        ],
        [
            "Покупатель",
            "ФИО и телефон повторяются в каждой продаже",
            (
                "Хранится один раз в customers; "
                "в sales используется customer_id"
            ),
        ],
        [
            "Сотрудник",
            "ФИО повторяется в каждой продаже",
            (
                "Хранится один раз в employees; "
                "в sales используется employee_id"
            ),
        ],
        [
            "Книга",
            "ISBN, название и жанр повторяются в продажах",
            (
                "Хранится один раз в books; "
                "жанр связан по genre_id"
            ),
        ],
        [
            "Автор",
            "Имя автора повторяется в продажах",
            (
                "Хранится один раз в authors; "
                "связь с книгами через book_authors"
            ),
        ],
        [
            "Продажа",
            "Содержит все текстовые данные",
            (
                "Хранит номер, дату, количество, цену "
                "и внешние ключи"
            ),
        ],
        [
            "Связи",
            "Связи существуют только логически в тексте",
            (
                "Связи 1:M и M:M заданы внешними ключами "
                "и таблицей book_authors"
            ),
        ],
        [
            "Целостность",
            "Зависит от ручной корректности повторяющихся данных",
            (
                "Обеспечивается PRIMARY KEY, UNIQUE, "
                "FOREIGN KEY и CHECK"
            ),
        ],
    ]

    table_queries = {
        "customers": """
            SELECT
                customer_id AS "ID покупателя",
                full_name AS "Покупатель",
                phone AS "Телефон"
            FROM customers
            ORDER BY customer_id
        """,
        "employees": """
            SELECT
                employee_id AS "ID сотрудника",
                full_name AS "Сотрудник"
            FROM employees
            ORDER BY employee_id
        """,
        "genres": """
            SELECT
                genre_id AS "ID жанра",
                name AS "Жанр"
            FROM genres
            ORDER BY genre_id
        """,
        "authors": """
            SELECT
                author_id AS "ID автора",
                full_name AS "Автор"
            FROM authors
            ORDER BY author_id
        """,
        "books": """
            SELECT
                b.book_id AS "ID книги",
                b.isbn AS "ISBN",
                b.title AS "Книга",
                b.genre_id AS "ID жанра",
                g.name AS "Жанр"
            FROM books b
            JOIN genres g ON g.genre_id = b.genre_id
            ORDER BY b.book_id
        """,
        "book_authors": """
            SELECT
                ba.book_id AS "ID книги",
                b.title AS "Книга",
                ba.author_id AS "ID автора",
                a.full_name AS "Автор"
            FROM book_authors ba
            JOIN books b ON b.book_id = ba.book_id
            JOIN authors a ON a.author_id = ba.author_id
            ORDER BY ba.book_id, ba.author_id
        """,
        "sales": """
            SELECT
                sale_id AS "ID продажи",
                sale_number AS "Номер продажи",
                sold_at AS "Дата",
                customer_id AS "ID покупателя",
                employee_id AS "ID сотрудника",
                book_id AS "ID книги",
                quantity AS "Количество",
                unit_price AS "Цена, ₽"
            FROM sales
            ORDER BY sale_id
        """,
    }

    REPORTS.mkdir(exist_ok=True)

    book = Workbook()
    book.remove(book.active)

    add_sheet(
        book,
        "Сравнение",
        comparison_headers,
        comparison_rows,
        "Сравнение: ненормализованная и нормализованная БД",
    )

    with psycopg.connect(pg_dsn()) as conn:
        with conn.cursor() as cur:
            for table_name, query in table_queries.items():
                cur.execute(query)

                headers = [
                    column.name
                    for column in cur.description
                ]

                rows = cur.fetchall()

                add_sheet(
                    book,
                    table_name,
                    headers,
                    rows,
                    f"Нормализованная PostgreSQL: {table_name}",
                )

    xlsx_path = REPORTS / "02_normalized_structure.xlsx"
    book.save(xlsx_path)

    return 8, xlsx_path


REPORT_QUERY = """
SELECT
    s.sold_at AS "Дата",
    COUNT(*) AS "Количество продаж",
    SUM(s.quantity) AS "Продано экземпляров",
    SUM(s.quantity * s.unit_price) AS "Выручка, ₽"
FROM sales s
GROUP BY s.sold_at
ORDER BY s.sold_at
"""


BEST_EMPLOYEE_QUERY = """
SELECT
    e.full_name AS "Лучший сотрудник",
    COUNT(s.sale_id) AS "Количество продаж",
    SUM(s.quantity) AS "Продано экземпляров",
    SUM(s.quantity * s.unit_price) AS "Выручка, ₽"
FROM sales s
JOIN employees e ON e.employee_id = s.employee_id
GROUP BY e.employee_id, e.full_name
ORDER BY
    SUM(s.quantity * s.unit_price) DESC,
    e.full_name
LIMIT 1
"""


LARGEST_SALE_QUERY = """
WITH ranked_sales AS (
    SELECT
        s.sold_at AS "Дата",
        s.sale_number AS "Номер продажи",
        c.full_name AS "Покупатель",
        c.phone AS "Телефон покупателя",
        e.full_name AS "Сотрудник",
        b.title AS "Книга",
        s.quantity AS "Количество",
        s.unit_price AS "Цена, ₽",
        s.quantity * s.unit_price AS "Сумма продажи, ₽",
        ROW_NUMBER() OVER (
            PARTITION BY s.sold_at
            ORDER BY
                s.quantity * s.unit_price DESC,
                s.sale_number
        ) AS rank_in_day
    FROM sales s
    JOIN customers c ON c.customer_id = s.customer_id
    JOIN employees e ON e.employee_id = s.employee_id
    JOIN books b ON b.book_id = s.book_id
)
SELECT
    "Дата",
    "Номер продажи",
    "Покупатель",
    "Телефон покупателя",
    "Сотрудник",
    "Книга",
    "Количество",
    "Цена, ₽",
    "Сумма продажи, ₽"
FROM ranked_sales
WHERE rank_in_day = 1
ORDER BY "Дата"
"""


def export_query(
    query: str,
    filename: str,
    sheet_name: str,
    title: str,
):
    """Выполняет SQL-запрос и экспортирует его в CSV и XLSX."""
    with psycopg.connect(pg_dsn()) as conn:
        with conn.cursor() as cur:
            cur.execute(query)

            headers = [
                column.name
                for column in cur.description
            ]

            rows = cur.fetchall()

    csv_path = REPORTS / f"{filename}.csv"
    xlsx_path = REPORTS / f"{filename}.xlsx"

    write_csv(csv_path, headers, rows)

    render_xlsx(
        csv_path,
        xlsx_path,
        sheet_name,
        title,
    )

    return len(rows), csv_path, xlsx_path


def export_report():
    """Формирует отчёт по продажам за каждый день."""
    return export_query(
        REPORT_QUERY,
        "03_sales_by_day",
        "Продажи по дням",
        "Шаг 3. Отчёт: продажи по дням",
    )


def export_best_employee():
    """
    Формирует отчёт по наиболее эффективному сотруднику.

    Критерий эффективности — максимальная суммарная выручка.
    """
    return export_query(
        BEST_EMPLOYEE_QUERY,
        "04_best_employee",
        "Лучший сотрудник",
        "Шаг 4. Самый эффективный сотрудник по выручке",
    )


def export_largest_sale_by_day():
    """
    Формирует отчёт о крупнейшей продаже каждого дня.

    Критерий — максимальная сумма:
    количество * цена.
    """
    return export_query(
        LARGEST_SALE_QUERY,
        "05_largest_sale_by_day",
        "Крупнейшие продажи",
        "Шаг 5. Самая большая продажа каждого дня",
    )
